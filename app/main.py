import json
import os
import re
from contextlib import suppress
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pypdf import PdfReader
from sqlalchemy import func
from sqlalchemy.orm import Session

from .ai_services import extract_text_from_image, generate_questions, grade_submission_with_ai
from .database import Base, engine, get_db
from .models import (
    AuthToken,
    Correction,
    Course,
    Evaluation,
    EvaluationQuestion,
    ExamSchedule,
    MaterialSource,
    Student,
    StudentEvaluation,
    Submission,
    Teacher,
)
from .security import create_token, hash_password, hash_token, verify_password
from .schemas import (
    AuthResponse,
    CorrectionCreate,
    CorrectionRead,
    CourseCreate,
    CourseRead,
    EvaluationCreate,
    EvaluationQuestionCreate,
    EvaluationRead,
    ExamScheduleCreate,
    ExamScheduleRead,
    MaterialSourceRead,
    ProgressStats,
    StudentCreate,
    StudentEvaluationRead,
    StudentRead,
    SubmissionCreate,
    SubmissionRead,
    TeacherLogin,
    TeacherRead,
    TeacherRegister,
)


def _teacher_or_401(teacher: Teacher | None) -> Teacher:
    if not teacher:
        raise HTTPException(status_code=401, detail="Docente no autenticado")
    return teacher

app = FastAPI(
    title="Asisdoc API",
    version="0.2.0",
    description="Asistente docente con autenticacion, IA y OCR.",
)


@app.middleware("http")
async def strip_api_prefix(request, call_next):
    """
    Permite servir la API detr?s de Firebase Hosting con rewrite /api/**.
    Si llega /api/*, se remapea internamente a /*.
    """
    path = request.scope.get("path", "")
    if path == "/api":
        request.scope["path"] = "/"
    elif path.startswith("/api/"):
        request.scope["path"] = path[4:]
    return await call_next(request)

Base.metadata.create_all(bind=engine)
UPLOAD_DIR = Path(os.getenv("UPLOAD_DIR", "uploads"))
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
MATERIALS_DIR = Path(os.getenv("MATERIALS_DIR", "materials"))
MATERIALS_DIR.mkdir(parents=True, exist_ok=True)
WEB_DIR = Path(__file__).resolve().parent.parent / "web"
app.mount("/web", StaticFiles(directory=WEB_DIR), name="web")


def get_current_teacher(
    authorization: str | None = Header(default=None),
    db: Session = Depends(get_db),
) -> Teacher:
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(status_code=401, detail="Token de acceso requerido")
    token = authorization.split(" ", 1)[1].strip()
    token_digest = hash_token(token)
    record = (
        db.query(AuthToken)
        .filter(
            AuthToken.token_hash == token_digest,
            AuthToken.revoked_at.is_(None),
        )
        .first()
    )
    if not record:
        raise HTTPException(status_code=401, detail="Token inválido")
    if record.expires_at < datetime.utcnow():
        raise HTTPException(status_code=401, detail="Token expirado")
    teacher = db.get(Teacher, record.teacher_id)
    if not teacher:
        raise HTTPException(status_code=401, detail="Docente no encontrado")
    return teacher


@app.post("/auth/register", response_model=TeacherRead)
def register_teacher(payload: TeacherRegister, db: Session = Depends(get_db)):
    email = payload.email.strip().lower()
    exists = db.query(Teacher).filter(Teacher.email == email).first()
    if exists:
        raise HTTPException(status_code=409, detail="El correo ya está registrado")
    teacher = Teacher(
        full_name=payload.full_name.strip(),
        email=email,
        password_hash=hash_password(payload.password),
    )
    db.add(teacher)
    db.commit()
    db.refresh(teacher)
    return teacher


@app.post("/auth/login", response_model=AuthResponse)
def login_teacher(payload: TeacherLogin, db: Session = Depends(get_db)):
    email = payload.email.strip().lower()
    teacher = db.query(Teacher).filter(Teacher.email == email).first()
    if not teacher or not verify_password(payload.password, teacher.password_hash):
        raise HTTPException(status_code=401, detail="Credenciales inválidas")

    plain_token, token_digest, expires_at = create_token()
    token = AuthToken(
        teacher_id=teacher.id,
        token_hash=token_digest,
        expires_at=expires_at,
    )
    db.add(token)
    db.commit()
    db.refresh(teacher)
    return AuthResponse(access_token=plain_token, expires_at=expires_at, teacher=teacher)


@app.post("/auth/logout")
def logout_teacher(
    authorization: str | None = Header(default=None),
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(status_code=401, detail="Token de acceso requerido")
    token = authorization.split(" ", 1)[1].strip()
    token_digest = hash_token(token)
    record = db.query(AuthToken).filter(AuthToken.token_hash == token_digest).first()
    if record and not record.revoked_at:
        record.revoked_at = datetime.utcnow()
        db.commit()
    return {"ok": True}


@app.get("/health")
def healthcheck():
    return {"status": "ok"}


@app.get("/")
def serve_web():
    index_path = WEB_DIR / "index.html"
    if index_path.exists():
        return FileResponse(
            index_path,
            headers={"Cache-Control": "no-store, no-cache, must-revalidate"},
        )
    raise HTTPException(status_code=404, detail="Interfaz web no disponible")


@app.get("/demo")
def serve_web_demo():
    index_path = WEB_DIR / "index.html"
    if index_path.exists():
        return FileResponse(
            index_path,
            headers={"Cache-Control": "no-store, no-cache, must-revalidate"},
        )
    raise HTTPException(status_code=404, detail="Interfaz demo no disponible")


@app.post("/courses", response_model=CourseRead)
def create_course(
    payload: CourseCreate,
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    exists = db.query(Course).filter(Course.name == payload.name).first()
    if exists:
        raise HTTPException(status_code=409, detail="El curso ya existe")
    course = Course(name=payload.name, description=payload.description)
    db.add(course)
    db.commit()
    db.refresh(course)
    return course


@app.get("/courses", response_model=list[CourseRead])
def list_courses(
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    return db.query(Course).order_by(Course.id.desc()).all()


@app.post("/students", response_model=StudentRead)
def create_student(
    payload: StudentCreate,
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    course = db.get(Course, payload.course_id)
    if not course:
        raise HTTPException(status_code=404, detail="Curso no encontrado")

    student = Student(full_name=payload.full_name, course_id=payload.course_id)
    db.add(student)
    db.commit()
    db.refresh(student)
    return student


@app.get("/students", response_model=list[StudentRead])
def list_students(
    course_id: Optional[int] = None,
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    query = db.query(Student)
    if course_id:
        query = query.filter(Student.course_id == course_id)
    return query.order_by(Student.id.desc()).all()


def _extract_text_from_uploaded_material(file_path: Path, extension: str) -> str:
    ext = extension.lower()
    if ext == ".pdf":
        with suppress(Exception):
            reader = PdfReader(str(file_path))
            pages = [page.extract_text() or "" for page in reader.pages]
            text = "\n".join(pages).strip()
            if text:
                return text[:30000]
    if ext == ".docx":
        with suppress(Exception):
            document = Document(str(file_path))
            text = "\n".join(p.text for p in document.paragraphs).strip()
            if text:
                return text[:30000]
    if ext in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"}:
        text = extract_text_from_image(file_path).strip()
        if text:
            return text[:30000]
    return ""


@app.post("/materials", response_model=MaterialSourceRead)
def upload_material(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    teacher = _teacher_or_401(_)
    allowed = {".pdf", ".docx", ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"}
    extension = Path(file.filename or "").suffix.lower()
    if extension not in allowed:
        raise HTTPException(
            status_code=400,
            detail="Formato no soportado. Permitidos: PDF, DOCX, PNG, JPG, JPEG, WEBP, BMP, TIFF.",
        )

    timestamp = int(datetime.utcnow().timestamp() * 1000)
    base_name = Path(file.filename or "material").stem.replace(" ", "_")
    safe_name = "".join(ch for ch in base_name if ch.isalnum() or ch in {"_", "-"})
    if not safe_name:
        safe_name = "material"
    filename = f"{safe_name}_{timestamp}{extension}"
    saved_path = MATERIALS_DIR / filename

    content = file.file.read()
    if not content:
        raise HTTPException(status_code=400, detail="El archivo está vacío.")
    with saved_path.open("wb") as out:
        out.write(content)

    extracted_text = _extract_text_from_uploaded_material(saved_path, extension)
    material = MaterialSource(
        teacher_id=teacher.id,
        original_filename=file.filename or filename,
        extension=extension,
        stored_path=str(saved_path),
        extracted_text=extracted_text,
    )
    db.add(material)
    db.commit()
    db.refresh(material)
    return material


@app.get("/materials", response_model=list[MaterialSourceRead])
def list_materials(
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    teacher = _teacher_or_401(_)
    return (
        db.query(MaterialSource)
        .filter(MaterialSource.teacher_id == teacher.id)
        .order_by(MaterialSource.id.desc())
        .all()
    )


@app.post("/evaluations", response_model=EvaluationRead)
def create_evaluation(
    payload: EvaluationCreate,
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    teacher = _teacher_or_401(_)
    course = db.get(Course, payload.course_id)
    if not course:
        raise HTTPException(status_code=404, detail="Curso no encontrado")

    material_text = payload.material_text
    if payload.material_source_id is not None:
        material = db.get(MaterialSource, payload.material_source_id)
        if not material or material.teacher_id != teacher.id:
            raise HTTPException(status_code=404, detail="Material no encontrado")
        material_text = material.extracted_text or material_text

    if payload.strict_material_only and not (material_text or "").strip():
        raise HTTPException(
            status_code=400,
            detail="No hay texto de material para restringir la generación.",
        )

    use_internal_knowledge = (
        payload.use_internal_knowledge and not payload.strict_material_only
    )
    material_available = bool((material_text or "").strip())
    force_material_fallback = payload.strict_material_only or payload.force_material_fallback
    # Pol?tica solicitada: si no hay material y no responde IA, no crear evaluaci?n vac?a.
    # Solo se permite respaldo autom?tico cuando existe material cargado.
    allow_fallback = material_available

    evaluation = Evaluation(
        title=payload.title,
        course_id=payload.course_id,
        evaluation_type=payload.evaluation_type,
        difficulty=payload.difficulty,
        material_text=material_text,
        material_source_id=payload.material_source_id,
        strict_material_only=payload.strict_material_only,
        use_internal_knowledge=use_internal_knowledge,
    )
    db.add(evaluation)
    db.flush()

    effective_topic = (payload.topic or "").strip() or payload.title

    try:
        generated_questions, generation_provider, used_fallback = generate_questions(
            topic=effective_topic,
            evaluation_type=payload.evaluation_type,
            difficulty=payload.difficulty,
            count=payload.question_count,
            material_text=material_text,
            restrict_to_material=payload.strict_material_only,
            allow_fallback=allow_fallback,
            force_material_fallback=force_material_fallback,
        )
    except RuntimeError as exc:
        db.rollback()
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    for question, answer in generated_questions:
        db.add(
            EvaluationQuestion(
                evaluation_id=evaluation.id, question_text=question, expected_answer=answer
            )
        )

    db.commit()
    db.refresh(evaluation)
    # Metadatos de trazabilidad para que el frontend informe IA real vs respaldo.
    evaluation.generation_provider = generation_provider
    evaluation.used_fallback = used_fallback
    return evaluation


@app.get("/evaluations", response_model=list[EvaluationRead])
def list_evaluations(
    course_id: Optional[int] = None,
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    query = db.query(Evaluation)
    if course_id:
        query = query.filter(Evaluation.course_id == course_id)
    return query.order_by(Evaluation.id.desc()).all()


@app.post("/evaluations/{evaluation_id}/questions", response_model=EvaluationRead)
def add_question(
    evaluation_id: int,
    payload: EvaluationQuestionCreate,
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    evaluation = db.get(Evaluation, evaluation_id)
    if not evaluation:
        raise HTTPException(status_code=404, detail="Evaluación no encontrada")
    db.add(
        EvaluationQuestion(
            evaluation_id=evaluation_id,
            question_text=payload.question_text,
            expected_answer=payload.expected_answer,
        )
    )
    db.commit()
    db.refresh(evaluation)
    return evaluation


@app.post("/schedules", response_model=ExamScheduleRead)
def create_schedule(
    payload: ExamScheduleCreate,
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    evaluation = db.get(Evaluation, payload.evaluation_id)
    if not evaluation:
        raise HTTPException(status_code=404, detail="Evaluación no encontrada")
    schedule = ExamSchedule(
        evaluation_id=payload.evaluation_id,
        scheduled_for=payload.scheduled_for,
        notes=payload.notes,
    )
    db.add(schedule)
    db.commit()
    db.refresh(schedule)
    return schedule


@app.get("/schedules", response_model=list[ExamScheduleRead])
def list_schedules(
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    return db.query(ExamSchedule).order_by(ExamSchedule.scheduled_for.asc()).all()


@app.post("/submissions", response_model=SubmissionRead)
def create_submission(
    payload: SubmissionCreate,
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    student = db.get(Student, payload.student_id)
    evaluation = db.get(Evaluation, payload.evaluation_id)
    if not student or not evaluation:
        raise HTTPException(
            status_code=404, detail="Alumno o evaluación no encontrada"
        )

    submission = Submission(
        student_id=payload.student_id,
        evaluation_id=payload.evaluation_id,
        image_reference=payload.image_reference,
        raw_text=payload.raw_text,
    )
    db.add(submission)
    db.commit()
    db.refresh(submission)
    return submission


@app.post("/submissions/photo", response_model=SubmissionRead)
def create_submission_from_photo(
    evaluation_id: int = Form(...),
    student_id: int = Form(...),
    image_file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    student = db.get(Student, student_id)
    evaluation = db.get(Evaluation, evaluation_id)
    if not student or not evaluation:
        raise HTTPException(
            status_code=404, detail="Alumno o evaluación no encontrada"
        )

    extension = Path(image_file.filename or "submission.jpg").suffix.lower()
    if extension not in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"}:
        raise HTTPException(status_code=400, detail="Formato de imagen no soportado")

    filename = f"sub_{student_id}_{evaluation_id}_{int(datetime.utcnow().timestamp())}{extension}"
    file_path = UPLOAD_DIR / filename
    with file_path.open("wb") as out:
        out.write(image_file.file.read())

    extracted_text = extract_text_from_image(file_path)
    submission = Submission(
        student_id=student_id,
        evaluation_id=evaluation_id,
        image_reference=str(file_path),
        raw_text=extracted_text,
    )
    db.add(submission)
    db.commit()
    db.refresh(submission)
    return submission


@app.post("/submissions/{submission_id}/correct", response_model=CorrectionRead)
def correct_submission(
    submission_id: int,
    payload: CorrectionCreate,
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    submission = db.get(Submission, submission_id)
    if not submission:
        raise HTTPException(status_code=404, detail="Entrega no encontrada")

    raw_text = (submission.raw_text or "").strip()
    evaluation = db.get(Evaluation, submission.evaluation_id)
    evaluation_questions = (
        db.query(EvaluationQuestion)
        .filter(EvaluationQuestion.evaluation_id == submission.evaluation_id)
        .order_by(EvaluationQuestion.id.asc())
        .all()
    )
    score, feedback, details = grade_submission_with_ai(
        raw_text=raw_text,
        criteria=payload.criteria,
        max_score=payload.max_score,
        use_llm=payload.use_llm,
        evaluation_type=evaluation.evaluation_type if evaluation else None,
        question_texts=[q.question_text for q in evaluation_questions],
    )
    details_json = json.dumps(details, ensure_ascii=False)
    full_feedback = f"{feedback}\n\nDETALLE_PREGUNTAS_JSON={details_json}"

    correction = Correction(
        submission_id=submission.id,
        criteria=payload.criteria,
        score=score,
        max_score=payload.max_score,
        feedback=full_feedback,
    )
    db.add(correction)

    student_eval = db.query(StudentEvaluation).filter(
        StudentEvaluation.student_id == submission.student_id,
        StudentEvaluation.evaluation_id == submission.evaluation_id,
    ).first()

    if student_eval:
        student_eval.score = score
        student_eval.max_score = payload.max_score
        student_eval.feedback = full_feedback
        student_eval.updated_at = datetime.utcnow()
    else:
        db.add(
            StudentEvaluation(
                student_id=submission.student_id,
                evaluation_id=submission.evaluation_id,
                score=score,
                max_score=payload.max_score,
                feedback=full_feedback,
            )
        )

    db.commit()
    db.refresh(correction)
    correction.detailed_feedback = feedback
    return correction


@app.get("/students/{student_id}/grades", response_model=list[StudentEvaluationRead])
def student_grades(
    student_id: int,
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    student = db.get(Student, student_id)
    if not student:
        raise HTTPException(status_code=404, detail="Alumno no encontrado")
    return (
        db.query(StudentEvaluation)
        .filter(StudentEvaluation.student_id == student_id)
        .order_by(StudentEvaluation.updated_at.desc())
        .all()
    )


@app.get("/students/{student_id}/progress", response_model=ProgressStats)
def student_progress(
    student_id: int,
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    student = db.get(Student, student_id)
    if not student:
        raise HTTPException(status_code=404, detail="Alumno no encontrado")

    rows = (
        db.query(StudentEvaluation.score, StudentEvaluation.max_score)
        .filter(StudentEvaluation.student_id == student_id)
        .all()
    )
    if not rows:
        return ProgressStats(
            student_id=student_id,
            evaluations_count=0,
            average_score=0.0,
            average_percentage=0.0,
            last_score=0.0,
        )

    evaluations_count = len(rows)
    avg_score = round(sum(r[0] for r in rows) / evaluations_count, 2)
    avg_percentage = round(
        sum((r[0] / r[1] * 100) if r[1] else 0 for r in rows) / evaluations_count, 2
    )
    last = (
        db.query(StudentEvaluation.score)
        .filter(StudentEvaluation.student_id == student_id)
        .order_by(StudentEvaluation.updated_at.desc())
        .first()
    )
    last_score = last[0] if last else 0.0

    return ProgressStats(
        student_id=student_id,
        evaluations_count=evaluations_count,
        average_score=avg_score,
        average_percentage=avg_percentage,
        last_score=last_score,
    )


def _humanize_evaluation_type(evaluation_type: str) -> str:
    mapping = {
        "multiple_choice": "Opcion multiple",
        "true_false": "Verdadero/Falso",
        "matching": "Relacionar conceptos",
        "open_answer": "Respuesta abierta",
    }
    return mapping.get(evaluation_type, evaluation_type)


def _parse_multiple_choice_question(question_text: str) -> tuple[str, dict[str, str]]:
    option_pattern = re.compile(r"^([A-D])[\)\.\-:]\s*(.+)$", flags=re.IGNORECASE)
    lines = [line.strip() for line in question_text.splitlines() if line.strip()]
    options: dict[str, str] = {}
    stem_lines: list[str] = []
    for line in lines:
        match = option_pattern.match(line)
        if match:
            options[match.group(1).upper()] = match.group(2).strip()
        else:
            stem_lines.append(line)

    stem = " ".join(stem_lines).strip() or question_text.strip()
    return stem, options


def _clean_export_stem(stem: str) -> str:
    cleaned = stem.strip()
    cleaned = re.sub(r"^\[[^\]]+\]\s*", "", cleaned)
    cleaned = re.sub(r"^\([^)]+\)\s*", "", cleaned)
    cleaned = re.sub(r"^Pregunta\s*\d+\s*:\s*", "", cleaned, flags=re.IGNORECASE)
    return cleaned.strip() or stem.strip()


def _parse_matching_question(
    question_text: str,
) -> tuple[str, list[tuple[str, str]], list[str], list[str]]:
    lines = [line.strip() for line in question_text.splitlines() if line.strip()]
    prompt_lines: list[str] = []
    pairs: list[tuple[str, str]] = []
    left_values: list[str] = []
    right_values: list[str] = []
    for line in lines:
        if "|" in line:
            left, right = [part.strip() for part in line.split("|", 1)]
            if left and right:
                pairs.append((left, right))
                left_values.append(left)
                right_values.append(right)
            continue
        prompt_lines.append(line)

    prompt = " ".join(prompt_lines).strip()
    prompt = _clean_export_stem(prompt) if prompt else "Relaciona cada elemento de la columna A con la columna B."
    return prompt, pairs, left_values, right_values


def _split_submission_answers(raw_text: str) -> list[str]:
    normalized = (raw_text or "").strip()
    if not normalized:
        return []
    lines = [line.strip() for line in normalized.splitlines() if line.strip()]
    answers: list[str] = []
    numbered_pattern = re.compile(r"^\s*(\d+)[\)\.\:\-]\s*(.+)$")
    for line in lines:
        match = numbered_pattern.match(line)
        if match:
            answers.append(match.group(2).strip())
            continue
        answers.append(line)
    return answers


def _build_feedback_items(
    answers: list[str], criteria: str, max_score: float
) -> tuple[list[dict], float]:
    keywords = [k.strip() for k in criteria.split(",") if k.strip()]
    total_questions = max(1, len(answers))
    points_per_question = round(max_score / total_questions, 2)
    result: list[dict] = []
    total_score = 0.0

    for idx, answer in enumerate(answers, start=1):
        normalized_answer = (answer or "").strip()
        lowered = normalized_answer.lower()
        if not lowered:
            result.append(
                {
                    "question": idx,
                    "status": "x",
                    "label": "X",
                    "score": 0.0,
                    "max_score": points_per_question,
                    "comment": "Sin respuesta. Marcada con X.",
                }
            )
            continue

        matched = sum(1 for kw in keywords if kw.lower() in lowered) if keywords else 0
        ratio = (matched / len(keywords)) if keywords else 0.6
        if ratio >= 0.75:
            score = points_per_question
            total_score += score
            result.append(
                {
                    "question": idx,
                    "status": "muy_bien",
                    "label": "Muy bien",
                    "score": round(score, 2),
                    "max_score": points_per_question,
                    "comment": f"Respuesta {idx}: Muy bien. Cumple con lo esperado.",
                }
            )
            continue

        if ratio >= 0.35:
            score = round(points_per_question * 0.6, 2)
            total_score += score
            missing = [kw for kw in keywords if kw.lower() not in lowered][:2]
            missing_hint = (
                f" Podr?as agregar: {', '.join(missing)}."
                if missing
                else " Podr?as ampliar con m?s precisi?n conceptual."
            )
            result.append(
                {
                    "question": idx,
                    "status": "mejorar",
                    "label": "Bien, mejorar",
                    "score": score,
                    "max_score": points_per_question,
                    "comment": (
                        f"Respuesta {idx}: Bien, pero incompleta."
                        f"{missing_hint}"
                    ),
                }
            )
            continue

        result.append(
            {
                "question": idx,
                "status": "x",
                "label": "X",
                "score": 0.0,
                "max_score": points_per_question,
                "comment": (
                    f"Respuesta {idx}: Incorrecta o fuera de criterio."
                    " Marcada con X."
                ),
            }
        )

    return result, round(min(max_score, total_score), 2)


def _build_teacher_feedback_text(items: list[dict], total_score: float, max_score: float) -> str:
    lines = [
        "Devolucion docente por pregunta:",
        *(f"- {item['comment']} ({item['score']}/{item['max_score']})" for item in items),
        f"Puntaje total: {round(total_score, 2)}/{max_score}",
    ]
    return "\n".join(lines)


def _extract_details_from_feedback(feedback_text: str) -> list[dict]:
    marker = "\n\nDETALLE_PREGUNTAS_JSON="
    if marker not in (feedback_text or ""):
        return []
    _, raw = feedback_text.split(marker, 1)
    try:
        parsed = json.loads(raw.strip())
        if isinstance(parsed, list):
            return parsed
    except Exception:
        return []
    return []


def _strip_feedback_details(feedback_text: str) -> str:
    marker = "\n\nDETALLE_PREGUNTAS_JSON="
    if marker not in (feedback_text or ""):
        return (feedback_text or "").strip()
    return (feedback_text.split(marker, 1)[0] or "").strip()


@app.get("/evaluations/{evaluation_id}/export-docx")
def export_evaluation_docx(
    evaluation_id: int,
    school_header: str = "Institucion Educativa",
    teacher_name: str = "Docente",
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    evaluation = db.get(Evaluation, evaluation_id)
    if not evaluation:
        raise HTTPException(status_code=404, detail="Evaluacion no encontrada")

    document = Document()
    title = document.add_heading(school_header, level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph(f"Evaluacion: {evaluation.title}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True

    meta = document.add_paragraph()
    meta.add_run("Docente: ").bold = True
    meta.add_run(teacher_name)
    meta.add_run("    |    ")
    meta.add_run("Tipo: ").bold = True
    meta.add_run(_humanize_evaluation_type(evaluation.evaluation_type))
    meta.add_run("    |    ")
    meta.add_run("Dificultad: ").bold = True
    meta.add_run(evaluation.difficulty.capitalize())

    document.add_paragraph("Nombre del alumno: ___________________________________________")
    document.add_paragraph("Curso: _______________________________    Fecha: ____ / ____ / ______")

    if evaluation.evaluation_type == "multiple_choice":
        document.add_paragraph(
            "Instrucciones: marque solo una opcion por pregunta con una X dentro de la casilla."
        )
    elif evaluation.evaluation_type == "true_false":
        document.add_paragraph(
            "Instrucciones: marque si cada afirmacion es Verdadera o Falsa."
        )
    elif evaluation.evaluation_type == "open_answer":
        document.add_paragraph(
            "Instrucciones: responda con claridad y fundamente cuando corresponda."
        )
    else:
        document.add_paragraph("Instrucciones: complete cada consigna segun lo solicitado.")

    document.add_paragraph("")

    questions = (
        db.query(EvaluationQuestion)
        .filter(EvaluationQuestion.evaluation_id == evaluation_id)
        .order_by(EvaluationQuestion.id.asc())
        .all()
    )
    for idx, q in enumerate(questions, start=1):
        if evaluation.evaluation_type == "multiple_choice":
            stem, options = _parse_multiple_choice_question(q.question_text)
            prompt = document.add_paragraph(f"{idx}. {_clean_export_stem(stem)}")
            prompt.runs[0].bold = True
            for letter in ("A", "B", "C", "D"):
                option_text = options.get(letter, "________________________________")
                document.add_paragraph(f"   [ ] {letter}) {option_text}")
            document.add_paragraph("")
            continue

        if evaluation.evaluation_type == "true_false":
            prompt = document.add_paragraph(f"{idx}. {_clean_export_stem(q.question_text)}")
            prompt.runs[0].bold = True
            document.add_paragraph("   [ ] Verdadero      [ ] Falso")
            document.add_paragraph("")
            continue

        if evaluation.evaluation_type == "matching":
            prompt, pairs, left_values, right_values = _parse_matching_question(q.question_text)
            question_paragraph = document.add_paragraph(f"{idx}. {prompt}")
            question_paragraph.runs[0].bold = True
            table = document.add_table(rows=max(2, len(left_values) + 1), cols=2)
            table.style = "Table Grid"
            table.cell(0, 0).text = "Columna A"
            table.cell(0, 1).text = "Columna B"
            for row_idx, value in enumerate(left_values, start=1):
                table.cell(row_idx, 0).text = f"{row_idx}) {value}"
            right_labels = [chr(65 + i) for i in range(len(right_values))]
            for row_idx, value in enumerate(right_values, start=1):
                label = right_labels[row_idx - 1] if row_idx - 1 < len(right_labels) else str(row_idx)
                table.cell(row_idx, 1).text = f"{label}) {value}"
            if pairs:
                document.add_paragraph(
                    "Indicacion: une con flechas los elementos de la Columna A con la Columna B."
                )
            document.add_paragraph("")
            continue

        document.add_paragraph(f"{idx}. {_clean_export_stem(q.question_text)}")
        if evaluation.evaluation_type == "open_answer":
            document.add_paragraph("Respuesta: _________________________________________________")
            document.add_paragraph("____________________________________________________________")
            document.add_paragraph("")

    file_stream = BytesIO()
    document.save(file_stream)
    file_stream.seek(0)

    filename = f"evaluacion_{evaluation_id}.docx"
    return StreamingResponse(
        file_stream,
        media_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/corrections/{correction_id}/export-docx")
def export_correction_docx(
    correction_id: int,
    school_header: str = "Institucion Educativa",
    teacher_name: str = "Docente",
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    correction = db.get(Correction, correction_id)
    if not correction:
        raise HTTPException(status_code=404, detail="Correccion no encontrada")
    submission = db.get(Submission, correction.submission_id)
    if not submission:
        raise HTTPException(status_code=404, detail="Entrega no encontrada")
    student = db.get(Student, submission.student_id)
    evaluation = db.get(Evaluation, submission.evaluation_id)
    if not student or not evaluation:
        raise HTTPException(status_code=404, detail="Datos relacionados incompletos")

    feedback_items = _extract_details_from_feedback(correction.feedback)
    if not feedback_items:
        answers = _split_submission_answers(submission.raw_text or "")
        feedback_items, _ = _build_feedback_items(
            answers=answers,
            criteria=correction.criteria,
            max_score=correction.max_score,
            question_count=max(1, len(answers)),
        )

    document = Document()
    heading = document.add_heading(school_header, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(f"Informe de correccion: {evaluation.title}")
    document.add_paragraph(f"Docente: {teacher_name}")
    document.add_paragraph(f"Alumno: {student.full_name}")
    document.add_paragraph(f"Puntaje final: {correction.score}/{correction.max_score}")
    document.add_paragraph("")
    document.add_paragraph("Detalle por pregunta:")

    for item in feedback_items:
        p = document.add_paragraph(
            f"Pregunta {item['question']}: {item['label']} - {item['score']}/{item['max_score']}"
        )
        p.runs[0].bold = True
        document.add_paragraph(item["comment"])

    document.add_paragraph("")
    document.add_paragraph("Resumen docente:")
    document.add_paragraph(_strip_feedback_details(correction.feedback))
    document.add_paragraph("")
    document.add_paragraph("Criterios utilizados:")
    document.add_paragraph(correction.criteria)

    file_stream = BytesIO()
    document.save(file_stream)
    file_stream.seek(0)
    filename = f"correccion_{correction_id}.docx"
    return StreamingResponse(
        file_stream,
        media_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/stats/courses/{course_id}")
def course_stats(
    course_id: int,
    db: Session = Depends(get_db),
    _: Teacher = Depends(get_current_teacher),
):
    course = db.get(Course, course_id)
    if not course:
        raise HTTPException(status_code=404, detail="Curso no encontrado")

    student_count = db.query(func.count(Student.id)).filter(Student.course_id == course_id).scalar()
    evaluation_count = (
        db.query(func.count(Evaluation.id))
        .filter(Evaluation.course_id == course_id)
        .scalar()
    )
    avg_percentage = (
        db.query(func.avg((StudentEvaluation.score / StudentEvaluation.max_score) * 100))
        .join(Student, Student.id == StudentEvaluation.student_id)
        .filter(Student.course_id == course_id)
        .scalar()
    )

    return {
        "course_id": course_id,
        "students": student_count or 0,
        "evaluations": evaluation_count or 0,
        "average_percentage": round(float(avg_percentage or 0.0), 2),
    }
